# สร้าง docx โครงการฝึกอบรมฉบับ กบค. สำหรับหลักสูตร Data Skills with Python (2 วัน)
# โครงหัวข้อตามต้นแบบ "โครงการฝึกอบรม Data Science สำหรับนักมาตรวิทยา 2.docx"
# เนื้อหาเป็นกลาง ไม่ผูกมาตรวิทยา · pandas (ไม่ใช่ Polars) · ไม่มี ML (ADR-0001, ADR-0002)
# ใช้: uv run python scripts/make_project_docx.py  (หรือ python3 ที่มี python-docx)
"""Generate the กบค. project-proposal docx for the 2-day Data Skills course."""

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

FONT = "TH Sarabun New"


def _style_run(run, size=16.0, bold=False):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = rpr.makeelement(qn("w:rFonts"), {})
        rpr.append(rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs"):
        rfonts.set(qn(attr), FONT)
    szcs = rpr.find(qn("w:szCs"))
    if szcs is None:
        szcs = rpr.makeelement(qn("w:szCs"), {})
        rpr.append(szcs)
    szcs.set(qn("w:val"), str(int(size * 2)))


def add_par(doc, text="", size=16.0, bold=False, align=None, space_after=0):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(0)
    if align is not None:
        p.alignment = align
    run = p.add_run(text)
    _style_run(run, size=size, bold=bold)
    return p


def add_heading(doc, text, size=16.0):
    return add_par(doc, text, size=size, bold=True, space_after=4)


def add_bullet(doc, text, indent=0.75):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.left_indent = Cm(indent)
    run = p.add_run("• " + text)
    _style_run(run)
    return p


def add_numbered(doc, text, indent=1.4):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.left_indent = Cm(indent)
    run = p.add_run(text)
    _style_run(run)
    return p


def add_cell_par(cell, text, bold=False, first=False):
    p = cell.paragraphs[0] if first and not cell.paragraphs[0].runs else cell.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(text)
    _style_run(run, size=16.0 if bold else 15.0, bold=bold)
    return p


def add_timetable(doc, title, rows):
    add_heading(doc, title, size=16.0)
    table = doc.add_table(rows=1 + len(rows), cols=3)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    widths = (Cm(3.2), Cm(11.8), Cm(4.0))
    header = ("ช่วงเวลา", "หัวข้อ / กิจกรรม", "รูปแบบ")
    for col, text in enumerate(header):
        cell = table.rows[0].cells[col]
        add_cell_par(cell, text, bold=True, first=True)
    for r, row in enumerate(rows, start=1):
        for c, value in enumerate(row):
            cell = table.rows[r].cells[c]
            if isinstance(value, str):
                values = [(value, False)]
            else:
                values = value
            for i, (text, bold) in enumerate(values):
                add_cell_par(cell, text, bold=bold, first=(i == 0))
    for row in table.rows:
        for c, cell in enumerate(row.cells):
            cell.width = widths[c]
    add_par(doc, "", space_after=6)


def build(out_path):
    doc = Document()
    for section in doc.sections:
        section.page_width = Cm(21.0)
        section.page_height = Cm(29.7)
        section.top_margin = Cm(2.0)
        section.bottom_margin = Cm(2.0)
        section.left_margin = Cm(2.0)
        section.right_margin = Cm(2.0)

    center = WD_ALIGN_PARAGRAPH.CENTER

    # ---- หัวกระดาษ ----
    add_par(doc, "โครงการฝึกอบรมเชิงปฏิบัติการ", size=18.0, bold=True, align=center)
    add_par(doc, "หลักสูตร Data Skills with Python", size=18.0, bold=True, align=center)
    add_par(doc, "จัดการข้อมูลจริง ตั้งแต่ไฟล์ดิบถึงรายงาน", size=18.0, bold=True, align=center)
    add_par(doc, "", space_after=6)
    add_par(doc, "หลักสูตรระยะเวลา 2 วัน", size=16.0, align=center)
    add_par(doc, "กลุ่มงานบริหารทรัพยากรบุคคล (กบค.) | สถาบันมาตรวิทยาแห่งชาติ", size=16.0, align=center)
    add_par(doc, "", space_after=6)

    add_par(doc, "ชื่อโครงการ", bold=True, space_after=2)
    add_par(
        doc,
        "โครงการฝึกอบรมเชิงปฏิบัติการ หลักสูตร Data Skills with Python — "
        "จัดการข้อมูลจริง ตั้งแต่ไฟล์ดิบถึงรายงาน "
        "เพื่อยกระดับทักษะการจัดการและวิเคราะห์ข้อมูลของบุคลากร",
        space_after=6,
    )
    add_par(doc, "หน่วยงานผู้จัด", bold=True, space_after=2)
    add_par(doc, "กลุ่มงานบริหารทรัพยากรบุคคล (กบค.) สถาบันมาตรวิทยาแห่งชาติ", space_after=6)
    add_par(doc, "กลุ่มเป้าหมาย", bold=True, space_after=2)
    add_par(
        doc,
        "บุคลากรทุกสายงาน จำนวน 20–25 คน "
        "(มีพื้นฐานการใช้คอมพิวเตอร์เบื้องต้น ไม่จำเป็นต้องมีพื้นฐานการเขียนโปรแกรม)",
        space_after=6,
    )
    add_par(doc, "ระยะเวลาและสถานที่", bold=True, space_after=2)
    add_par(
        doc,
        "2 วัน (รวม 11 ชั่วโมง 30 นาที) เวลา 09:00 - 16:00 น. "
        "ณ ห้องฝึกอบรม สถาบันมาตรวิทยาแห่งชาติ",
        space_after=6,
    )
    add_par(doc, "รูปแบบการจัดอบรม", bold=True, space_after=2)
    add_par(
        doc,
        "บรรยายปูพื้นฐานควบคู่กับการทำ Guided Hands-on / Code-Along "
        "โดยวิทยากรจัดเตรียมชุดข้อมูลตัวอย่างและแม่แบบโค้ด (Template Notebooks) ไว้ล่วงหน้า "
        "ผู้เข้าอบรมคัดลอกตัวอย่างคำสั่ง รันตามทีละขั้นตอน (Step-by-step execution) "
        "ทดลองปรับค่าพารามิเตอร์ และฟังคำอธิบายการแปลผลลัพธ์อย่างเป็นระบบ "
        "เพื่อลดอุปสรรคเรื่องข้อผิดพลาดทางไวยากรณ์ (Syntax errors) และมุ่งเน้นความเข้าใจเชิงตรรกะ",
        space_after=6,
    )

    # ---- 1. หลักการและเหตุผล ----
    add_heading(doc, "1. หลักการและเหตุผล")
    add_par(
        doc,
        "ปัจจุบันองค์กรทุกประเภททำงานอยู่กับข้อมูลในปริมาณที่เพิ่มขึ้นทุกปี "
        "ไม่ว่าจะเป็นไฟล์รายงานประจำเดือน ข้อมูลจากระบบงานต่าง ๆ แบบสำรวจ "
        "หรือข้อมูลจากเครื่องมือและเซนเซอร์ การทำงานด้วยสเปรดชีตแบบเดิมเริ่มมีข้อจำกัด "
        "ทั้งด้านขนาดข้อมูล การทำงานซ้ำ ๆ ด้วยมือ และความผิดพลาดจากการคัดลอกวาง "
        "ที่ตรวจสอบย้อนกลับได้ยาก",
        space_after=6,
    )
    add_par(
        doc,
        "ภาษา Python ควบคู่กับไลบรารี pandas คือเครื่องมือจัดการข้อมูลแบบตาราง "
        "ที่ได้รับความนิยมและใช้กันแพร่หลายที่สุดในวงการข้อมูลทั่วโลก "
        "เป็นซอฟต์แวร์ฟรีและโอเพนซอร์ส มีชุมชนผู้ใช้ขนาดใหญ่ และมีสื่อการเรียนรู้มากมาย "
        "ทำให้เหมาะเป็นจุดเริ่มต้นของการยกระดับทักษะข้อมูลของบุคลากรทุกสายงาน "
        "โดยสามารถโหลดไฟล์ข้อมูลจริง จัดการข้อมูลที่ไม่สะอาด สรุปด้วยสถิติ "
        "สร้างกราฟ และพยากรณ์แนวโน้มอย่างง่ายได้ครบภายในหลักสูตรเดียว",
        space_after=6,
    )
    add_par(
        doc,
        "โครงการนี้ออกแบบสำหรับผู้เริ่มต้นโดยเฉพาะ ผู้เรียนไม่จำเป็นต้องมีพื้นฐานการเขียนโปรแกรม "
        "เน้นการเรียนรู้แบบ Code-Along คือรันโค้ดต้นแบบที่วิทยากรเตรียมไว้ "
        "พร้อมอธิบายเหตุผลว่าแต่ละหัวข้อจำเป็นต่องานประจำวันอย่างไร "
        "และปิดท้ายด้วย Mini Project ที่ผู้เรียนได้ลงมือทำครบวงจรกับข้อมูลของตนเอง "
        "ทักษะที่ได้จากหลักสูตรนี้ยังเป็นรากฐานสำหรับหลักสูตรการวิเคราะห์ข้อมูลขั้นสูงที่จะจัดต่อไป",
        space_after=6,
    )

    # ---- 2. วัตถุประสงค์ ----
    add_heading(doc, "2. วัตถุประสงค์")
    objectives = (
        "เข้าใจรูปแบบไฟล์ข้อมูลที่พบในการทำงาน (CSV, Excel, JSON) "
        "ขีดจำกัดของการทำงานด้วยสเปรดชีต และใช้งาน Notebook "
        "พร้อมเขียน Python พื้นฐาน (ตัวแปร ลูป ฟังก์ชัน) แบบรันและแก้ไขได้",
        "จัดการตารางข้อมูลด้วย pandas ได้ ทั้งการโหลดไฟล์ CSV/Excel "
        "การกรอง เรียงลำดับ สรุปข้อมูลเป็นกลุ่ม และการรวมหลายไฟล์เป็นตารางเดียว",
        "ตรวจจับและจัดการข้อมูลไม่สะอาด (ค่าว่าง สะกดคลาดเคลื่อน วันที่แปลก ไฟล์ซ้ำซ้อน) "
        "ให้เป็นตารางสะอาด พร้อมสรุปข้อมูลด้วยสถิติ "
        "และนำเสนอด้วยกราฟที่เลือกให้ตรงกับคำถาม",
        "สร้างการพยากรณ์แนวโน้มอย่างง่าย (trend line / moving average) ได้ "
        "พร้อมประเมินได้ว่าการพยากรณ์เชื่อถือได้แค่ไหน "
        "และประยุกต์ใช้กระบวนการทั้งหมดกับข้อมูลของตนเองใน Mini Project",
    )
    for i, obj in enumerate(objectives, start=1):
        add_numbered(doc, f"{i}. {obj}", indent=0.75)
    add_par(doc, "", space_after=6)

    # ---- 3. กำหนดการ วันที่ 1 ----
    day1 = [
        (
            "09:00 - 09:30",
            [("ลงทะเบียน / ชี้แจงวัตถุประสงค์ / ทดสอบก่อนอบรม (Pre-test)", False)],
            "กิจกรรมเปิด",
        ),
        (
            "09:30 - 10:30",
            [
                ("Session 1: ข้อมูลรอบตัวเรา", True),
                ("• ข้อมูลอยู่รอบตัวเรา: รูปแบบไฟล์ที่พบในงานประจำวัน (CSV, Excel, JSON) "
                 "และความแตกต่างของแต่ละรูปแบบ", False),
                ("• ขีดจำกัดของสเปรดชีต: ขนาดไฟล์ งานซ้ำที่ต้องทำด้วยมือ "
                 "ข้อผิดพลาดจากการคัดลอกวาง และการตรวจสอบย้อนกลับที่ยาก", False),
                ("• ทำไมต้องเรียนเรื่องนี้: งานเอกสารและรายงานส่วนใหญ่เริ่มจาก \"ไฟล์ข้อมูล\" "
                 "เมื่อเข้าใจโครงสร้างของไฟล์และขีดจำกัดของเครื่องมือเดิม "
                 "จะรู้ว่าเมื่อไหร่ควรเปลี่ยนไปใช้เครื่องมือที่ทำซ้ำได้และตรวจสอบได้", False),
            ],
            "Lecture",
        ),
        ("10:30 - 10:45", "พักรับประทานอาหารว่าง", "พัก"),
        (
            "10:45 - 12:00",
            [
                ("Session 2: Python จากศูนย์", True),
                ("• แนะนำสภาพแวดล้อมการทำงาน (Jupyter Notebook) และการรันคำสั่งทีละเซลล์", False),
                ("• Python พื้นฐานแบบ Code-Along: ตัวแปร ข้อความ ตัวเลข ลูป และฟังก์ชันสำเร็จรูป", False),
                ("• Workshop 1 (Guided Code-Along):", True),
                ("1. คัดลอกและรันโค้ดตัวอย่างการอ่านไฟล์ข้อมูลและแสดงผลบนจอ", False),
                ("2. แก้ไขค่าในโค้ดและสังเกตผลลัพธ์ทันที", False),
                ("3. ฟังคำอธิบายว่าโค้ดแต่ละบรรทัดทำอะไร โดยไม่ต้องจำไวยากรณ์", False),
            ],
            "Guided Hands-on (Code-Along)",
        ),
        ("12:00 - 12:45", "พักรับประทานอาหารกลางวัน", "พัก"),
        (
            "12:45 - 14:00",
            [
                ("Session 3: จัดการตารางข้อมูลด้วย pandas", True),
                ("• ทำความรู้จัก pandas: โครงสร้างตาราง (DataFrame) ที่เป็นมาตรฐานของงานวิเคราะห์ข้อมูล", False),
                ("• โหลดไฟล์ CSV/Excel เข้า Python กรอง เรียงลำดับ และสรุปข้อมูลเป็นกลุ่ม (group by)", False),
                ("• Workshop 2 (Guided Code-Along):", True),
                ("1. รันโค้ดโหลดไฟล์ข้อมูลตัวอย่างจริง (CSV/Excel)", False),
                ("2. รันคำสั่งกรอง เรียง และสรุปข้อมูลตามกลุ่มตามโค้ดต้นแบบ", False),
                ("3. ทดลองเปลี่ยนเงื่อนไขการกรองและสังเกตการเปลี่ยนแปลงของผลลัพธ์", False),
            ],
            "Guided Hands-on (Code-Along)",
        ),
        ("14:00 - 14:15", "พักรับประทานอาหารว่าง", "พัก"),
        (
            "14:15 - 15:45",
            [
                ("Session 4: ข้อมูลไม่สะอาด", True),
                ("• ปัญหาข้อมูลไม่สะอาดที่พบบ่อย: ค่าว่าง ชื่อสะกดคลาดเคลื่อน "
                 "วันที่รูปแบบไม่ตรง ไฟล์ซ้ำซ้อนจากหลายแหล่ง", False),
                ("• ทำไมต้องเรียนเรื่องนี้: สถิติและกราฟที่เชื่อถือได้ต้องเริ่มจากข้อมูลที่สะอาด "
                 "การรู้จักปัญหาและวิธีตรวจจับตั้งแต่ต้นช่วยประหยัดเวลา "
                 "และป้องกันข้อสรุปที่คลาดเคลื่อน", False),
                ("• Workshop 3 (Guided Code-Along):", True),
                ("1. รันโค้ดตรวจพบค่าว่างและค่าซ้ำซ้อนในไฟล์ตัวอย่าง", False),
                ("2. รันคำสั่งจัดการข้อมูลไม่สะอาดทีละขั้น (เติมค่า ลบซ้ำ แก้รูปแบบวันที่)", False),
                ("3. ปิดท้ายด้วยตารางสะอาดพร้อมใช้วิเคราะห์", False),
            ],
            "Guided Hands-on (Code-Along)",
        ),
        (
            "15:45 - 16:00",
            "สรุปประเด็นสำคัญประจำวัน ตอบข้อซักถาม และแนะนำเอกสารประกอบ",
            "สรุป",
        ),
    ]
    add_timetable(doc, "3. กำหนดการ วันที่ 1: จัดการข้อมูลเป็น — จากไฟล์ดิบสู่ตารางสะอาด", day1)

    # ---- 4. กำหนดการ วันที่ 2 ----
    day2 = [
        (
            "09:00 - 10:15",
            [
                ("Session 5: สถิติที่ใช้ได้จริง", True),
                ("• สถิติพรรณนาที่ใช้สื่อสาร: ค่าเฉลี่ย มัธยฐาน การกระจาย และ percentile", False),
                ("• \"เมื่อไหร่ค่าเฉลี่ยหลอกเรา\": ตัวอย่างข้อมูลที่ค่าเฉลี่ยนำทางผิด "
                 "และวิธีอ่านข้อมูลให้รอบคอบ", False),
                ("• Workshop 4 (Guided Code-Along):", True),
                ("1. รันโค้ดสรุปข้อมูลด้วยสถิติพรรณนาจากตารางที่จัดการแล้วในวันแรก", False),
                ("2. เปรียบเทียบค่าเฉลี่ยกับมัธยฐานและอ่านการกระจายของข้อมูล", False),
            ],
            "Lecture + Guided Hands-on",
        ),
        ("10:15 - 10:30", "พักรับประทานอาหารว่าง", "พัก"),
        (
            "10:30 - 11:45",
            [
                ("Session 6: กราฟที่สื่อสารได้", True),
                ("• เลือกชนิดกราฟให้ตรงคำถาม: การกระจาย (distribution) "
                 "การเปรียบเทียบ (comparison) ความสัมพันธ์ (relationship)", False),
                ("• พล็อตกราฟด้วย Matplotlib และหลักการจัดกราฟให้อ่านง่าย (ชื่อกราฟ แกน หน่วย)", False),
                ("• Workshop 5 (Guided Code-Along):", True),
                ("1. รันโค้ดวาดกราฟชนิดต่าง ๆ จากข้อมูลตัวอย่างจริง", False),
                ("2. เลือกกราฟให้ตรงกับคำถามที่ต้องการตอบ แล้วปรับรายละเอียดให้สื่อสารได้", False),
            ],
            "Guided Hands-on (Code-Along)",
        ),
        ("11:45 - 12:30", "พักรับประทานอาหารกลางวัน", "พัก"),
        (
            "12:30 - 13:45",
            [
                ("Session 7: พยากรณ์อย่างง่าย", True),
                ("• เส้นแนวโน้ม (trend line) และค่าเฉลี่ยเคลื่อนที่ (moving average) "
                 "สำหรับข้อมูลตามเวลา", False),
                ("• ความผันผวนตามวงรอบ (ฤดูกาล วันหยุด) และการอ่านแนวโน้มอย่างมีวิจารณญาณ", False),
                ("• กับดัก extrapolation: ทำไมการทำนายที่ยาวเกินช่วงข้อมูลจึงไม่น่าเชื่อถือ", False),
                ("• Workshop 6 (Guided Code-Along):", True),
                ("1. รันโค้ดสร้างเส้นแนวโน้มและค่าเฉลี่ยเคลื่อนที่จากข้อมูลตามเวลาจริง", False),
                ("2. ทดลองพยากรณ์ระยะสั้นและระยะยาว แล้วชี้จุดที่การพยากรณ์เริ่มไม่น่าเชื่อถือ", False),
            ],
            "Guided Hands-on (Code-Along)",
        ),
        ("13:45 - 14:00", "พักรับประทานอาหารว่าง", "พัก"),
        (
            "14:00 - 14:45",
            [
                ("Session 8: ทำงานอัตโนมัติ", True),
                ("• ประมวลผลไฟล์หลายไฟล์เป็นชุดด้วยลูป: รวมไฟล์รายเดือนเป็นตารางเดียวอัตโนมัติ", False),
                ("• สร้างรายงานจากแม่แบบ (template) อัตโนมัติ: ตารางสรุป + กราฟพร้อมส่งต่อ", False),
                ("• Workshop 7 (Guided Code-Along):", True),
                ("1. รันสคริปต์รวมไฟล์ข้อมูลหลายไฟล์และสร้างรายงานสรุปอัตโนมัติ", False),
            ],
            "Guided Hands-on (Code-Along)",
        ),
        (
            "14:45 - 15:40",
            [
                ("Session 9: Mini Project", True),
                ("• ผู้เรียนเลือกข้อมูลของตนเอง (หรือชุดข้อมูลสำรองที่จัดเตรียมไว้) "
                 "ทำครบวงจร: ไฟล์ดิบ → จัดการข้อมูล → กราฟ → พยากรณ์แนวโน้ม → รายงานสรุป", False),
                ("• วิทยากรและผู้ช่วยเดินให้คำแนะนำรายกลุ่ม (Coaching)", False),
            ],
            "Workshop (Group Work)",
        ),
        (
            "15:40 - 16:00",
            [
                ("• นำเสนอผลงาน Mini Project แบบรวดเร็ว (กลุ่มละ 2–3 นาที) "
                 "พร้อมตอบคำถาม \"ทำไมจึงเชื่อการพยากรณ์นี้ได้\"", False),
                ("• สไลด์ \"What's next\": ทักษะที่ได้วันนี้ (Python + ข้อมูลสะอาด + แนวโน้ม) "
                 "คือรากฐานของการวิเคราะห์ข้อมูลขั้นสูงที่จะต่อยอดต่อไป", False),
                ("• แบบทดสอบหลังการอบรม (Post-test) / ประเมินผล / ปิดการอบรม", False),
            ],
            "นำเสนอ + กิจกรรมปิด",
        ),
    ]
    add_timetable(doc, "4. กำหนดการ วันที่ 2: จากข้อมูล สู่ความเข้าใจ สู่การมองไปข้างหน้า", day2)

    # ---- 5. วิธีการจัดการเรียนรู้ ----
    add_heading(doc, "5. วิธีการจัดการเรียนรู้และการเตรียมความพร้อม")
    add_bullet(
        doc,
        "Guided Code-Along Strategy: วิทยากรแจกชุดเอกสาร Jupyter Notebook ที่มีโค้ดตัวอย่าง "
        "พร้อมคอมเมนต์อธิบายทุกบรรทัด ผู้เรียนคัดลอก/รันโค้ดตาม "
        "เพื่อเรียนรู้กระบวนการแบบเห็นผลลัพธ์ทันที โดยไม่ต้องเสียเวลากับการจำไวยากรณ์หรือแก้พิมพ์ผิด",
    )
    add_bullet(
        doc,
        "โฟกัสที่การแปลผลลัพธ์: ใช้เวลาส่วนใหญ่ไปกับการอธิบายว่าผลลัพธ์ ตาราง และกราฟ "
        "ที่โค้ดสร้างขึ้น มีความหมายอย่างไรต่อการตัดสินใจในงานประจำวัน",
    )
    add_bullet(doc, "ไลบรารีหลัก: ภาษา Python, pandas (จัดการตารางข้อมูล), Matplotlib (การสร้างกราฟ)")
    add_bullet(
        doc,
        "สภาพแวดล้อมพร้อมใช้แบบออฟไลน์: เครื่องผู้เรียนติดตั้งโปรแกรมและชุดข้อมูลสำเร็จล่วงหน้า "
        "ผ่านสคริปต์เตรียมเครื่อง การอบรมทั้งหมดไม่ต้องใช้อินเทอร์เน็ต",
    )
    add_par(doc, "", space_after=6)

    # ---- 6. การวัดและประเมินผล ----
    add_heading(doc, "6. การวัดและประเมินผล")
    add_bullet(doc, "แบบทดสอบวัดความรู้ความเข้าใจก่อนและหลังการอบรม (Pre-test / Post-test)")
    add_bullet(doc, "การทดสอบการรันและปรับแต่งค่าพารามิเตอร์ในไฟล์แบบฝึกหัดรายบุคคล")
    add_bullet(
        doc,
        "ผลงาน Mini Project รายกลุ่ม ที่แสดงกระบวนการทำงานครบวงจร "
        "และตอบได้ว่าการพยากรณ์เชื่อถือได้แค่ไหน",
    )
    add_bullet(doc, "แบบสอบถามประเมินความพึงพอใจและประสิทธิผลของหลักสูตร")
    add_par(doc, "", space_after=6)

    # ---- 7. ประโยชน์ที่คาดว่าจะได้รับ ----
    add_heading(doc, "7. ประโยชน์ที่คาดว่าจะได้รับ")
    add_bullet(
        doc,
        "บุคลากรใช้ Python และ pandas จัดการไฟล์ข้อมูลประจำวันที่มีขนาดใหญ่ "
        "หรือมาจากหลายไฟล์ได้ด้วยตนเอง ลดงานซ้ำที่ต้องทำด้วยมือ",
    )
    add_bullet(
        doc,
        "ผู้เรียนเลือกสถิติและชนิดกราฟได้ตรงคำถาม "
        "และสื่อสารข้อมูลกับผู้บริหารและเพื่อนร่วมงานได้ชัดเจนขึ้น",
    )
    add_bullet(
        doc,
        "ผู้เรียนสร้างการพยากรณ์แนวโน้มอย่างง่ายได้ "
        "พร้อมรู้เท่าทันขีดจำกัดของการพยากรณ์ (ไม่เชื่อการพยากรณ์เกินจริง)",
    )
    add_bullet(
        doc,
        "หน่วยงานมีต้นแบบโค้ดและแม่แบบรายงาน (Templates) "
        "ที่นำกลับไปปรับใช้กับงานของแต่ละสายงานได้ทันที",
    )
    add_bullet(
        doc,
        "บุคลากรมีรากฐานพร้อมสำหรับหลักสูตรการวิเคราะห์ข้อมูลขั้นสูงที่จะจัดต่อไป",
    )

    doc.save(out_path)
    print(f"saved: {out_path}")


if __name__ == "__main__":
    build("docs/proposal/โครงการฝึกอบรม Data Skills with Python (ฉบับ กบค.).docx")
